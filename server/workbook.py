"""Bounded workbook serializer. JSON data only; no arbitrary code/formulas/network."""
import sys,json,io,math
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter

data=json.loads(sys.stdin.buffer.read(250000));sheets=data['sheets']
assert 1<=len(sheets)<=3
book=Workbook();book.remove(book.active);book.properties.creator='Companion Agent';book.properties.title='用户任务表格'
for item in sheets:
 assert isinstance(item['name'],str) and 1<=len(item['name'])<=31
 rows=item['rows'];assert 1<=len(rows)<=101
 ws=book.create_sheet(item['name'])
 for ri,row in enumerate(rows,1):
  assert 1<=len(row)<=12
  for ci,value in enumerate(row,1):
   assert value is None or isinstance(value,(str,int,float,bool))
   if isinstance(value,str):assert len(value)<=500 and '\0' not in value
   if isinstance(value,float):assert math.isfinite(value)
   cell=ws.cell(ri,ci,value)
   # Spreadsheet formula injection is disabled: even a leading '=' stays text.
   if isinstance(value,str):cell.data_type='s'
   cell.font=Font(name='Microsoft YaHei',size=11,color='30261F',bold=ri==1)
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if ri==1:cell.fill=PatternFill('solid',fgColor='EAD5BC')
   elif ri%2==0:cell.fill=PatternFill('solid',fgColor='FAF6F0')
   if isinstance(value,float):cell.number_format='0.00'
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=25
 for column in ws.columns:
  n=max(len(str(c.value or '')) for c in column);ws.column_dimensions[column[0].column_letter].width=min(38,max(14,n*1.3+3))
 ws.sheet_view.showGridLines=False;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
out=io.BytesIO();book.save(out);raw=out.getvalue();assert len(raw)<2*1024*1024
check=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
assert check.sheetnames==[x['name'] for x in sheets]
for sheet,expected in zip(check.worksheets,sheets):
 actual=list(sheet.iter_rows(values_only=True));assert actual==[tuple(None if v=='' else v for v in row) for row in expected['rows']]
check.close();sys.stdout.buffer.write(raw)
